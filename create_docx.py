#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
아이캔노트 사용설명서 Word 문서 생성 스크립트
"""

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_document_style(doc):
    """문서 기본 스타일 설정"""
    # 기본 폰트 설정
    style = doc.styles['Normal']
    style.font.name = '맑은 고딕'
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

    # 제목 스타일 설정
    for i in range(1, 4):
        heading_style = doc.styles[f'Heading {i}']
        heading_style.font.name = '맑은 고딕'
        heading_style.font.bold = True
        heading_style._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

def add_heading_with_color(doc, text, level, color=None):
    """색상이 있는 제목 추가"""
    heading = doc.add_heading(text, level)
    if color:
        for run in heading.runs:
            run.font.color.rgb = color
    return heading

def create_icannote_manual():
    """아이캔노트 사용설명서 생성"""
    doc = Document()
    set_document_style(doc)

    # === 표지 ===
    title = doc.add_heading('아이캔노트 (ICanNote)', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph('사용 설명서')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(24)
    subtitle.runs[0].font.bold = True

    doc.add_paragraph()
    doc.add_paragraph()

    version_info = doc.add_paragraph('무료 전자 판서 프로그램')
    version_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    version_info.runs[0].font.size = Pt(14)

    doc.add_paragraph()

    website = doc.add_paragraph('공식 웹사이트: https://icannote.com')
    website.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()

    # === 목차 ===
    doc.add_heading('목차', 1)
    toc_items = [
        '1. 프로그램 소개',
        '2. 설치 방법',
        '3. 기본 인터페이스',
        '4. 주요 기능',
        '   4.1 필기 도구',
        '   4.2 도형 도구',
        '   4.3 선택 및 편집 도구',
        '   4.4 화면 조작',
        '   4.5 녹화 기능',
        '   4.6 파일 관리',
        '5. 교육 환경별 활용법',
        '   5.1 온라인 강의 녹화',
        '   5.2 실시간 원격수업',
        '   5.3 교실 대면수업',
        '   5.4 수학 수업',
        '   5.5 어학 수업',
        '6. 설정 방법',
        '7. 단축키 모음',
        '8. 자주 묻는 질문 (FAQ)',
        '9. 개발 스토리',
        '10. 저작권 및 라이센스',
    ]
    for item in toc_items:
        doc.add_paragraph(item)

    doc.add_page_break()

    # === 1. 프로그램 소개 ===
    doc.add_heading('1. 프로그램 소개', 1)

    intro_text = """아이캔노트(ICanNote)는 디지털 강의를 위해 제작된 무료 전자 판서 프로그램입니다. 컴퓨터를 사용하여 강의를 진행하시는 모든 분들이 프레젠테이션 화면에 필기하고, 도형을 그리며, 강조할 수 있는 강력한 도구입니다.

PDF, HWP, PPT, DOC, JPG 등 다양한 형식의 교재 파일을 불러와서 그 위에 직접 판서할 수 있으며, 빈 화이트보드 화면에서도 자유롭게 필기할 수 있습니다. 강의 내용을 음성과 함께 녹화하여 동영상으로 저장할 수도 있어, 온라인 강의 제작에 최적화되어 있습니다."""
    doc.add_paragraph(intro_text)

    doc.add_heading('주요 특징', 2)
    features = [
        '완전 무료: 개인, 기업, 기관, 단체 등 누구나 무상으로 사용 가능',
        '다양한 필기 도구: 연필, 형광펜, 도형 그리기 등 풍부한 판서 도구 제공',
        '강의 녹화: 화면과 음성을 함께 녹화하여 동영상 강의 제작 가능',
        '다양한 파일 지원: PDF, HWP, PPT, DOC, JPG 등 다양한 교재 파일 불러오기',
        '저장 및 공유: ICN 형식으로 저장하여 판서 내용 보존 및 수정 가능',
        '저작권 보장: 제작한 강의의 저작권은 사용자에게 귀속',
    ]
    for feature in features:
        doc.add_paragraph(f'• {feature}')

    doc.add_page_break()

    # === 2. 설치 방법 ===
    doc.add_heading('2. 설치 방법', 1)

    doc.add_heading('시스템 요구사항', 2)
    requirements = [
        '운영체제: Windows 7/8/10/11',
        '저장공간: 약 100MB 이상의 여유 공간',
        '권장: 타블렛 펜 또는 터치스크린 (마우스로도 사용 가능)',
    ]
    for req in requirements:
        doc.add_paragraph(f'• {req}')

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run('※ 참고: ').bold = True
    p.add_run('스마트폰, 스마트패드와 같은 안드로이드 또는 iOS는 지원되지 않습니다.')

    doc.add_heading('설치 절차', 2)
    install_steps = [
        '아이캔노트 공식 웹사이트(https://icannote.com)에 접속합니다.',
        '다운로드 버튼을 클릭하여 설치 파일을 다운로드합니다.',
        '다운로드 받은 압축 파일을 원하는 위치에 압축 해제합니다.',
        'setup.exe 파일을 더블클릭하여 실행합니다.',
        '설치 마법사의 안내에 따라 설치를 완료합니다.',
        '바탕화면에 생성된 아이캔노트 아이콘을 더블클릭하여 실행합니다.',
    ]
    for i, step in enumerate(install_steps, 1):
        doc.add_paragraph(f'{i}. {step}')

    doc.add_heading('ffmpeg 코덱 설치 (녹화 기능 사용 시)', 2)
    codec_text = """동영상 녹화 기능을 사용하기 위해서는 ffmpeg 코덱을 추가로 설치해야 합니다. 처음 녹화 버튼을 누르면 코덱이 설치되어 있지 않으면 설치 안내가 표시됩니다. 안내에 따라 설치하시면 됩니다."""
    doc.add_paragraph(codec_text)

    doc.add_page_break()

    # === 3. 기본 인터페이스 ===
    doc.add_heading('3. 기본 인터페이스', 1)

    interface_text = """아이캔노트의 화면은 크게 다음과 같은 영역으로 구성되어 있습니다."""
    doc.add_paragraph(interface_text)

    interface_areas = [
        ('상단 메뉴바', '파일 열기, 저장, 설정 등의 메뉴가 위치합니다.'),
        ('좌측 도구바', '필기 도구, 도형 도구, 선택 도구 등이 세로로 배열되어 있습니다.'),
        ('중앙 작업 영역', '실제로 판서가 이루어지는 메인 작업 공간입니다.'),
        ('우측 하단 컨트롤', '녹화 버튼, 페이지 이동 버튼 등이 위치합니다.'),
        ('하단 도구 옵션', '선택한 도구의 세부 옵션(색상, 두께 등)을 조절할 수 있습니다.'),
    ]
    for area_name, description in interface_areas:
        p = doc.add_paragraph()
        p.add_run(f'• {area_name}: ').bold = True
        p.add_run(description)

    doc.add_page_break()

    # === 4. 주요 기능 ===
    doc.add_heading('4. 주요 기능', 1)

    # 4.1 필기 도구
    doc.add_heading('4.1 필기 도구', 2)

    writing_tools = [
        ('자유선 (펜)', '화면에 자유롭게 글을 쓰거나 그림을 그릴 수 있습니다. 마우스나 타블렛 펜으로 자연스러운 필기가 가능합니다.'),
        ('형광펜', '텍스트나 중요한 부분을 강조할 때 사용합니다. 반투명한 색상으로 표시되어 아래 내용이 보입니다.'),
        ('직선 도구', '깔끔한 직선을 그릴 수 있습니다. Shift키를 누르면 수평/수직/45도 각도의 선을 그릴 수 있습니다.'),
        ('점 지우개', '필기한 내용을 부분적으로 지울 수 있습니다. 칠판 지우개처럼 문지르면 해당 부분만 지워집니다.'),
        ('획 지우개', '하나의 필기 획 전체를 한 번에 지웁니다.'),
    ]
    for tool_name, description in writing_tools:
        p = doc.add_paragraph()
        p.add_run(f'■ {tool_name}\n').bold = True
        p.add_run(description)

    p = doc.add_paragraph()
    p.add_run('※ 팁: ').bold = True
    p.add_run('지우개는 필기만 지우고, 배경이나 이미지는 지우지 않습니다. 도형이나 이미지를 지우려면 선택 도구로 선택 후 Delete키를 누르세요.')

    # 4.2 도형 도구
    doc.add_heading('4.2 도형 도구', 2)

    shape_tools = [
        ('사각형', '사각형을 그립니다. Shift키를 누르면 정사각형이 됩니다.'),
        ('원/타원', '원이나 타원을 그립니다. Shift키를 누르면 정원이 됩니다.'),
        ('삼각형', '삼각형을 그립니다. Shift키를 누르면 정삼각형이 됩니다.'),
        ('다각형', '육각형 등 다양한 각의 다각형을 그릴 수 있습니다.'),
        ('화살표', '방향을 표시하는 화살표를 그립니다.'),
        ('스티커', '가리고 싶은 부분을 가릴 때 사용합니다. 선택 도구로 클릭하면 다시 보이게 할 수 있습니다.'),
    ]
    for tool_name, description in shape_tools:
        p = doc.add_paragraph()
        p.add_run(f'■ {tool_name}\n').bold = True
        p.add_run(description)

    # 4.3 선택 및 편집 도구
    doc.add_heading('4.3 선택 및 편집 도구', 2)

    select_tools = [
        ('선택 도구 (올가미)', '필기 내용, 그림, 도형 등을 선택합니다. 클릭으로 개별 선택하거나, 드래그로 범위를 지정하여 여러 개를 한꺼번에 선택할 수 있습니다.'),
        ('이동', '선택한 객체를 원하는 위치로 이동시킵니다.'),
        ('크기 조절', '선택한 객체의 크기를 조절합니다.'),
        ('회전', '이미지나 도형을 90도씩 시계 방향으로 회전시킵니다. 하단 툴박스의 회전 아이콘을 클릭합니다.'),
        ('삭제', '선택한 객체를 삭제합니다. Delete키 또는 하단의 X 아이콘을 클릭합니다.'),
    ]
    for tool_name, description in select_tools:
        p = doc.add_paragraph()
        p.add_run(f'■ {tool_name}\n').bold = True
        p.add_run(description)

    # 4.4 화면 조작
    doc.add_heading('4.4 화면 조작', 2)

    screen_tools = [
        ('손바닥 도구 (화면 이동)', '손바닥 모양을 선택하면 화면에 있는 문서를 이동시킬 수 있습니다. 화면을 확대했을 때 보이지 않는 부분으로 이동할 때 유용합니다.'),
        ('화면 확대', 'Ctrl + PgUp 또는 확대 버튼을 클릭합니다. 세부 내용을 자세히 보거나 정밀한 작업을 할 때 사용합니다.'),
        ('화면 축소', 'Ctrl + PgDn 또는 축소 버튼을 클릭합니다. 전체 페이지를 한눈에 보고 싶을 때 사용합니다.'),
        ('페이지 이동', '화면 하단의 페이지 이동 버튼을 사용하거나, 페이지 목록에서 원하는 페이지를 선택합니다.'),
    ]
    for tool_name, description in screen_tools:
        p = doc.add_paragraph()
        p.add_run(f'■ {tool_name}\n').bold = True
        p.add_run(description)

    # 4.5 녹화 기능
    doc.add_heading('4.5 녹화 기능', 2)

    recording_text = """아이캔노트의 강력한 기능 중 하나는 화면과 음성을 함께 녹화하여 동영상 강의를 제작할 수 있다는 것입니다."""
    doc.add_paragraph(recording_text)

    doc.add_heading('녹화 방법', 3)
    recording_steps = [
        '우측 하단의 톱니바퀴(설정)를 클릭합니다.',
        '녹화 설정 화면에서 마이크 입력장치와 시스템 사운드를 설정합니다.',
        '설정을 마친 후 녹화 시작(재생) 버튼을 클릭합니다.',
        '필요한 경우 녹화할 영역을 지정할 수 있습니다.',
        '강의가 끝나면 중지 버튼을 클릭합니다.',
        '녹화된 파일이 지정된 폴더에 저장됩니다.',
    ]
    for i, step in enumerate(recording_steps, 1):
        doc.add_paragraph(f'{i}. {step}')

    p = doc.add_paragraph()
    p.add_run('※ 참고: ').bold = True
    p.add_run('처음 녹화 시 ffmpeg 코덱 설치가 필요할 수 있습니다. 안내에 따라 설치하시면 됩니다.')

    # 4.6 파일 관리
    doc.add_heading('4.6 파일 관리', 2)

    file_text = """아이캔노트는 다양한 형식의 파일을 불러오고 저장할 수 있습니다."""
    doc.add_paragraph(file_text)

    doc.add_heading('지원 파일 형식', 3)
    file_formats = [
        ('불러오기 지원', 'PDF, HWP, PPT, DOC, DOCX, JPG, PNG, GIF 등'),
        ('저장 형식', 'ICN (아이캔노트 전용 형식) - 판서 내용 수정 가능'),
        ('내보내기', '이미지 파일, PDF 등으로 내보내기 가능'),
    ]
    for format_name, formats in file_formats:
        p = doc.add_paragraph()
        p.add_run(f'• {format_name}: ').bold = True
        p.add_run(formats)

    doc.add_heading('파일 저장하기', 3)
    save_text = """상단 메뉴의 [저장/인쇄 - 저장하기]를 선택하면 *.icn 형식으로 저장됩니다. ICN 파일에는 교재로 사용된 원본 파일의 내용과 판서한 내용이 함께 저장되어 있어, 나중에 불러와서 수정이 가능합니다."""
    doc.add_paragraph(save_text)

    doc.add_page_break()

    # === 5. 교육 환경별 활용법 ===
    doc.add_heading('5. 교육 환경별 활용법', 1)

    # 5.1 온라인 강의 녹화
    doc.add_heading('5.1 온라인 강의 녹화', 2)

    online_recording = """온라인 강의 콘텐츠 제작에 아이캔노트를 활용하는 방법입니다."""
    doc.add_paragraph(online_recording)

    online_tips = [
        '강의에 사용할 교재(PDF, PPT 등)를 미리 불러옵니다.',
        '녹화 전 마이크와 시스템 사운드 설정을 확인합니다.',
        '화면 해상도와 녹화 품질을 적절히 설정합니다.',
        '필기 색상은 배경과 대비되는 색을 선택합니다.',
        '중요한 내용은 형광펜이나 스티커로 강조합니다.',
        '녹화 중에도 자유롭게 페이지를 이동하며 강의합니다.',
        '녹화 완료 후 영상을 확인하고 필요시 재녹화합니다.',
    ]
    for tip in online_tips:
        doc.add_paragraph(f'• {tip}')

    # 5.2 실시간 원격수업
    doc.add_heading('5.2 실시간 원격수업', 2)

    realtime_text = """Zoom, Google Meet, MS Teams 등의 화상회의 도구와 함께 사용하여 실시간 원격수업을 진행할 수 있습니다."""
    doc.add_paragraph(realtime_text)

    realtime_tips = [
        '화상회의 프로그램의 화면 공유 기능으로 아이캔노트 화면을 공유합니다.',
        '학생들이 판서 내용을 잘 볼 수 있도록 펜 두께를 적절히 조절합니다.',
        '화면 공유 시 "컴퓨터 소리 포함" 옵션을 활성화합니다.',
        '수업 중 학생 질문에 즉시 필기로 답변할 수 있습니다.',
        '중요한 수업은 동시에 녹화하여 복습 자료로 제공합니다.',
    ]
    for tip in realtime_tips:
        doc.add_paragraph(f'• {tip}')

    # 5.3 교실 대면수업
    doc.add_heading('5.3 교실 대면수업', 2)

    classroom_text = """프로젝터나 대형 디스플레이와 연결하여 교실 대면수업에서도 활용할 수 있습니다."""
    doc.add_paragraph(classroom_text)

    classroom_tips = [
        '프로젝터/디스플레이에 컴퓨터 화면을 연결합니다.',
        '배경을 칠판 색상(청록색 그라데이션)으로 설정하면 익숙한 느낌을 줍니다.',
        '펜 두께를 크게 설정하여 뒷자리 학생도 볼 수 있게 합니다.',
        '타블렛 펜을 사용하면 더 자연스러운 필기가 가능합니다.',
        '수업 내용을 저장하여 학생들에게 공유할 수 있습니다.',
    ]
    for tip in classroom_tips:
        doc.add_paragraph(f'• {tip}')

    # 5.4 수학 수업
    doc.add_heading('5.4 수학 수업', 2)

    math_text = """수학 수업에서 아이캔노트를 특히 효과적으로 활용할 수 있습니다."""
    doc.add_paragraph(math_text)

    math_tips = [
        '격자 배경을 활성화하여 그래프와 도형을 정확하게 그립니다.',
        '도형 도구로 정확한 삼각형, 사각형, 원 등을 그립니다.',
        '수식을 단계별로 풀이하며 필기합니다.',
        '다른 색상을 사용하여 풀이 과정의 각 단계를 구분합니다.',
        '스티커 기능으로 정답을 가렸다가 공개하는 방식으로 활용합니다.',
        '문제 이미지를 불러와 그 위에 직접 풀이를 작성합니다.',
    ]
    for tip in math_tips:
        doc.add_paragraph(f'• {tip}')

    # 5.5 어학 수업
    doc.add_heading('5.5 어학 수업', 2)

    language_text = """외국어 수업에서도 아이캔노트를 다양하게 활용할 수 있습니다."""
    doc.add_paragraph(language_text)

    language_tips = [
        '텍스트가 있는 교재를 불러와 중요 어휘에 밑줄이나 형광펜을 사용합니다.',
        '단어의 발음 기호나 뜻을 필기로 추가합니다.',
        '문장 구조를 색상별로 구분하여 표시합니다.',
        '스티커로 정답을 가리고 학생들에게 맞추게 하는 퀴즈 형식으로 활용합니다.',
        '녹화 기능으로 발음과 함께 설명하는 영상을 제작합니다.',
        '대화문을 역할별로 다른 색상으로 표시합니다.',
    ]
    for tip in language_tips:
        doc.add_paragraph(f'• {tip}')

    doc.add_page_break()

    # === 6. 설정 방법 ===
    doc.add_heading('6. 설정 방법', 1)

    settings_intro = """상단 메뉴의 [설정]을 클릭하면 다양한 옵션을 조절할 수 있습니다."""
    doc.add_paragraph(settings_intro)

    # 배경 설정
    doc.add_heading('배경 설정', 2)
    background_settings = [
        '[설정 - 배경/마우스 커서]를 선택합니다.',
        '"페이지 배경"에서 원하는 배경 색상을 선택합니다.',
        '그라데이션을 선택하면 칠판과 같은 느낌의 배경이 적용됩니다.',
        '격자무늬 배경은 [페이지 선 그리기]에서 설정할 수 있습니다.',
    ]
    for setting in background_settings:
        doc.add_paragraph(f'• {setting}')

    # 펜 설정
    doc.add_heading('펜 설정', 2)
    pen_settings = [
        '펜 두께: 하단 도구바에서 원하는 두께를 선택합니다.',
        '펜 색상: 하단 도구바의 색상 팔레트에서 선택합니다.',
        '사용자 지정 색상: 숫자키 1~5에 자주 사용하는 색상을 지정할 수 있습니다.',
    ]
    for setting in pen_settings:
        doc.add_paragraph(f'• {setting}')

    # 녹화 설정
    doc.add_heading('녹화 설정', 2)
    recording_settings = [
        '우측 하단의 톱니바퀴 아이콘을 클릭합니다.',
        '마이크 입력장치를 선택합니다.',
        '시스템 사운드 포함 여부를 설정합니다.',
        '녹화 파일 저장 위치를 지정합니다.',
        '녹화 품질(해상도, 프레임레이트)을 설정합니다.',
    ]
    for setting in recording_settings:
        doc.add_paragraph(f'• {setting}')

    doc.add_page_break()

    # === 7. 단축키 모음 ===
    doc.add_heading('7. 단축키 모음', 1)

    # 테이블 생성
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'

    # 헤더
    header_cells = table.rows[0].cells
    header_cells[0].text = '단축키'
    header_cells[1].text = '기능'
    for cell in header_cells:
        cell.paragraphs[0].runs[0].bold = True

    # 단축키 목록
    shortcuts = [
        ('Ctrl + PgUp', '화면 확대'),
        ('Ctrl + PgDn', '화면 축소'),
        ('Ctrl + Z', '실행 취소'),
        ('Ctrl + Y', '다시 실행'),
        ('Ctrl + S', '저장하기'),
        ('Ctrl + O', '파일 열기'),
        ('Delete', '선택한 객체 삭제'),
        ('1 ~ 5', '사용자 지정 색상 전환'),
        ('Shift + 도형 그리기', '정비례 도형 (정사각형, 정원 등)'),
        ('Shift + 직선', '수평/수직/45도 직선'),
    ]

    for shortcut, function in shortcuts:
        row_cells = table.add_row().cells
        row_cells[0].text = shortcut
        row_cells[1].text = function

    doc.add_page_break()

    # === 8. 자주 묻는 질문 (FAQ) ===
    doc.add_heading('8. 자주 묻는 질문 (FAQ)', 1)

    faqs = [
        ('Q: 아이캔노트는 유료인가요?',
         'A: 아니요, 완전 무료입니다. 개인, 기업, 기관, 단체 등 누구나 무상으로 사용할 수 있습니다.'),

        ('Q: 맥(Mac)이나 스마트폰에서 사용할 수 있나요?',
         'A: 현재 Windows 운영체제에서만 사용 가능합니다. macOS, 안드로이드, iOS는 지원되지 않습니다.'),

        ('Q: 녹화가 안 됩니다. 어떻게 해야 하나요?',
         'A: ffmpeg 코덱이 설치되어 있는지 확인하세요. 녹화 버튼을 처음 누르면 코덱 설치 안내가 나타납니다. 또한 녹화 설정에서 마이크와 시스템 사운드가 올바르게 설정되어 있는지 확인하세요.'),

        ('Q: 지우개로 도형이 지워지지 않아요.',
         'A: 점 지우개는 자유선(펜 필기)만 지웁니다. 도형, 직선, 이미지를 지우려면 선택 도구(올가미)로 선택한 후 Delete키를 누르거나 하단의 X 아이콘을 클릭하세요.'),

        ('Q: 아이캔노트로 만든 강의 영상의 저작권은 누구에게 있나요?',
         'A: 아이캔노트로 제작한 강의 영상이나 출력물의 저작권은 사용자에게 있습니다. 유튜브나 다른 매체에서 수익을 창출해도 됩니다.'),

        ('Q: HWP 파일을 불러올 수 없어요.',
         'A: HWP 파일을 불러오기 위해서는 한글 프로그램이 설치되어 있어야 합니다. 또는 HWP 파일을 PDF로 변환한 후 불러오시면 됩니다.'),

        ('Q: 저장한 ICN 파일을 다른 컴퓨터에서 열 수 있나요?',
         'A: 네, 다른 컴퓨터에 아이캔노트가 설치되어 있다면 ICN 파일을 열어 수정하거나 이어서 작업할 수 있습니다.'),

        ('Q: 화면이 너무 작아서 필기하기 어려워요.',
         'A: Ctrl + PgUp으로 화면을 확대하세요. 확대 후 손바닥 도구로 화면을 이동할 수 있습니다.'),

        ('Q: 배경을 칠판처럼 바꾸고 싶어요.',
         'A: [설정 - 배경/마우스 커서]에서 "페이지 배경"의 그라데이션 옵션을 선택하면 칠판 색상으로 변경됩니다.'),

        ('Q: 타블렛 없이도 사용할 수 있나요?',
         'A: 네, 마우스로도 사용 가능합니다. 다만 자연스러운 필기를 위해서는 타블렛 펜 사용을 권장합니다.'),
    ]

    for question, answer in faqs:
        p = doc.add_paragraph()
        p.add_run(question).bold = True
        doc.add_paragraph(answer)
        doc.add_paragraph()

    doc.add_page_break()

    # === 9. 개발 스토리 ===
    doc.add_heading('9. 개발 스토리', 1)

    doc.add_heading('개발 의도', 2)
    dev_intent = """아이캔노트는 강사와 학생의 거리감을 줄이고, 보다 이해하기 쉬운 강의 제작에 도움을 드릴 목적으로 개발되었습니다.

온라인 교육이 확산되면서 많은 선생님들께서 강의 녹화와 전자 판서에 어려움을 겪고 계셨습니다. 기존의 상용 프로그램들은 가격이 비싸거나 기능이 복잡하여 접근하기 어려웠습니다.

이러한 문제를 해결하고자, 현직 교육자이자 수학 강사인 이상열 대표를 비롯한 소수의 선생님과 개발자들이 모여 아이캔노트를 만들게 되었습니다. 누구나 쉽게 사용할 수 있고, 무료로 제공되는 전자 판서 프로그램을 목표로 개발하였습니다."""
    doc.add_paragraph(dev_intent)

    doc.add_heading('열정으로 만들어가는 프로젝트', 2)
    passion_text = """아이캔노트는 극소수의 개발자들이 대가 없이, 오직 열정만으로 작업하고 있는 프로젝트입니다.

정규 업무 외 시간을 쪼개어 프로그램을 개선하고, 사용자분들의 피드백을 반영하여 업데이트를 진행하고 있습니다. 어떠한 광고도 없이, 오로지 교육 현장에서 더 나은 강의가 이루어지기를 바라는 마음으로 개발을 이어가고 있습니다.

사용해 주시는 모든 선생님들과 학생들에게 감사드리며, 앞으로도 더 좋은 프로그램이 될 수 있도록 노력하겠습니다."""
    doc.add_paragraph(passion_text)

    doc.add_heading('후원 안내', 2)

    support_text = """아이캔노트의 지속적인 개발과 운영에는 서버 비용, 개발 도구 라이센스 등 비용이 발생합니다.

저희 개발팀에게 작은 응원의 마음을 전해주시면 큰 힘이 됩니다. 매달 햄버거 하나만 사주셔도 개발자들에게는 더 열심히 개발할 수 있는 원동력이 됩니다."""
    doc.add_paragraph(support_text)

    p = doc.add_paragraph()
    p.add_run('\n☕ 후원하기: ').bold = True
    p.add_run('공식 웹사이트(https://icannote.com)에서 후원 방법을 확인하실 수 있습니다.\n')

    gratitude = doc.add_paragraph('\n후원해 주시는 모든 분들께 진심으로 감사드립니다. 여러분의 응원이 아이캔노트를 더욱 발전시키는 원동력입니다.')
    gratitude.runs[0].italic = True

    doc.add_page_break()

    # === 10. 저작권 및 라이센스 ===
    doc.add_heading('10. 저작권 및 라이센스', 1)

    license_text = """아이캔노트와 아이캔스크린은 개인, 기업, 기관, 단체 등 사용자(처)의 구분과 관계없이 무상으로 사용이 가능한 소프트웨어입니다."""
    doc.add_paragraph(license_text)

    doc.add_heading('저작권 안내', 2)
    copyright_items = [
        '아이캔노트/스크린을 이용하여 강의한 동영상이나 출력물에 대한 저작권은 사용자에게 있습니다.',
        '유튜브나 다른 매체를 이용해서 수익을 창출해도 됩니다. 수익은 사용자 본인이 모두 가져가시면 됩니다.',
        '상업적 용도로 자유롭게 사용하실 수 있습니다.',
        '별도의 라이센스 구매나 등록 절차가 필요 없습니다.',
    ]
    for item in copyright_items:
        doc.add_paragraph(f'• {item}')

    doc.add_paragraph()

    # 마무리
    final_text = doc.add_paragraph('\n본 설명서에 관한 문의사항은 공식 웹사이트를 통해 문의해 주시기 바랍니다.')

    website_final = doc.add_paragraph()
    website_final.add_run('\n공식 웹사이트: ').bold = True
    website_final.add_run('https://icannote.com')

    # 문서 저장
    doc.save('/home/user/icannote/아이캔노트_사용설명서.docx')
    print('아이캔노트_사용설명서.docx 파일이 생성되었습니다.')

if __name__ == '__main__':
    create_icannote_manual()
